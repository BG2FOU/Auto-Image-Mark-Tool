"""Preview CSV, TSV and XLSX edits against stable photo IDs before applying them."""

from __future__ import annotations

import csv
import zipfile
from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass, replace
from datetime import date, datetime
from io import StringIO
from pathlib import Path, PurePosixPath
from typing import Any
from uuid import UUID

from openpyxl import load_workbook  # type: ignore[import-untyped]

from aim_tool.domain import PhotoItem
from aim_tool.domain.dates import DateError, DateInput, parse_capture_date
from aim_tool.domain.validation import validate_coordinates, validate_watermark_fields

FIELDS = (
    "file_name",
    "category",
    "subject",
    "capture_date",
    "location_name",
    "latitude",
    "longitude",
)
HEADERS = {
    **{field: field for field in FIELDS},
    "文件名": "file_name",
    "类别": "category",
    "内容": "subject",
    "拍摄日期": "capture_date",
    "地点": "location_name",
    "纬度": "latitude",
    "经度": "longitude",
}


class TableImportError(ValueError):
    """A table cannot be parsed or safely applied."""


@dataclass(frozen=True)
class ImportedRow:
    line: int
    values: dict[str, str | date | datetime]


@dataclass(frozen=True)
class MatchedRow:
    photo_id: UUID
    row: ImportedRow


@dataclass(frozen=True)
class ImportPreview:
    matches: tuple[MatchedRow, ...]
    errors: tuple[str, ...]


def _headers(
    raw: tuple[Any, ...], column_mapping: Mapping[str, str] | None = None
) -> dict[int, str]:
    names = dict(HEADERS)
    for alias, field in (column_mapping or {}).items():
        if not alias.strip() or field not in FIELDS:
            raise TableImportError("Column mapping needs a nonempty header and a supported field")
        names[alias.strip()] = field
    result: dict[int, str] = {}
    for index, value in enumerate(raw):
        name = names.get(str(value).strip()) if value is not None else None
        if name is not None:
            if name in result.values():
                raise TableImportError(f"Duplicate mapped column: {name}")
            result[index] = name
    if "file_name" not in result.values():
        raise TableImportError("Table needs a file_name column")
    return result


def _rows(
    records: list[tuple[int, tuple[Any, ...]]], column_mapping: Mapping[str, str] | None = None
) -> tuple[ImportedRow, ...]:
    if not records:
        raise TableImportError("Table is empty")
    headers = _headers(records[0][1], column_mapping)
    rows: list[ImportedRow] = []
    for number, record in records[1:]:
        values: dict[str, str | date | datetime] = {}
        for index, name in headers.items():
            value = record[index] if index < len(record) else None
            if value is None or isinstance(value, str) and not value.strip():
                continue
            if name in {"file_name", "subject"} and not isinstance(value, str):
                raise TableImportError(
                    f"Row {number}: {name} must be text to preserve leading zeros"
                )
            if isinstance(value, datetime | date):
                if name != "capture_date":
                    raise TableImportError(f"Row {number}: date cell in {name} column")
                values[name] = value
            else:
                values[name] = str(value).strip()
        if values:
            rows.append(ImportedRow(number, values))
    return tuple(rows)


def read_pasted_tsv(
    text: str, *, column_mapping: Mapping[str, str] | None = None
) -> tuple[ImportedRow, ...]:
    """Parse actual clipboard tabs and quoted records before a user applies edits."""
    try:
        reader = csv.reader(
            StringIO(text.lstrip("\ufeff"), newline=""), delimiter="\t", strict=True
        )
        return _rows([(reader.line_num, tuple(record)) for record in reader], column_mapping)
    except csv.Error as error:
        raise TableImportError(f"Cannot read pasted TSV: {error}") from error


def read_table(
    path: Path,
    *,
    encoding: str = "utf-8-sig",
    sheet: str | None = None,
    column_mapping: Mapping[str, str] | None = None,
) -> tuple[ImportedRow, ...]:
    """Never infer a fallback encoding or evaluate spreadsheet formulas."""
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        if encoding not in {"utf-8-sig", "utf-8", "gb18030"}:
            raise TableImportError("Choose UTF-8 or explicit GB18030 encoding")
        if sheet is not None:
            raise TableImportError("Worksheet selection applies only to XLSX")
        try:
            with path.open("r", encoding=encoding, newline="") as source:
                delimiter = "\t" if suffix == ".tsv" else ","
                reader = csv.reader(source, delimiter=delimiter, strict=True)
                records = [(reader.line_num, tuple(record)) for record in reader]
        except (OSError, UnicodeError, csv.Error) as error:
            raise TableImportError(f"Cannot read table: {error}") from error
        return _rows(records, column_mapping)
    if suffix != ".xlsx":
        raise TableImportError("Only CSV, TSV and XLSX tables are supported")
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
        try:
            if workbook._external_links:
                raise TableImportError("External workbook links are not allowed")
            if sheet is not None and sheet not in workbook.sheetnames:
                raise TableImportError(f"Unknown worksheet: {sheet}")
            worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
            records = []
            for cells in worksheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise TableImportError(f"Formula cell on row {cells[0].row}; paste its value")
                records.append((cells[0].row, tuple(cell.value for cell in cells)))
            return _rows(records, column_mapping)
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        if isinstance(error, TableImportError):
            raise
        raise TableImportError(f"Cannot read XLSX: {error}") from error


def _file_key(value: str) -> str:
    cleaned = value.replace("\\", "/")
    parts = cleaned.split("/")
    if (
        cleaned.startswith("/")
        or any(part in {"", ".", ".."} for part in parts)
        or ":" in cleaned
        or "\x00" in cleaned
    ):
        raise TableImportError(f"Unsafe file_name: {value}")
    path = PurePosixPath(cleaned)
    return path.as_posix()


def _validate_values(row: ImportedRow, photo: PhotoItem) -> None:
    values = row.values
    category = values.get("category", photo.edits.get("category"))
    subject = values.get("subject", photo.edits.get("subject"))
    try:
        validate_watermark_fields(
            str(category) if category is not None else None,
            str(subject) if subject is not None else None,
        )
    except ValueError as error:
        raise TableImportError(f"Row {row.line}: {error}") from error
    if "capture_date" in values:
        try:
            parse_capture_date(values["capture_date"])
        except DateError as error:
            raise TableImportError(f"Row {row.line}: {error}") from error
    if ("latitude" in values) != ("longitude" in values):
        raise TableImportError(f"Row {row.line}: latitude and longitude must be paired")
    if "latitude" in values:
        try:
            validate_coordinates(float(str(values["latitude"])), float(str(values["longitude"])))
        except (ValueError, TypeError) as error:
            raise TableImportError(f"Row {row.line}: invalid WGS84 coordinates") from error


def preview_import(rows: tuple[ImportedRow, ...], photos: tuple[PhotoItem, ...]) -> ImportPreview:
    """Match relative paths exactly, or an unambiguous basename in this batch."""
    by_relative: dict[str, list[PhotoItem]] = {}
    by_name: dict[str, list[PhotoItem]] = {}
    for photo in photos:
        try:
            relative = photo.source.relative_to(photo.import_root).as_posix()
        except ValueError as error:
            raise TableImportError("Photo is outside its import root") from error
        by_relative.setdefault(relative, []).append(photo)
        by_name.setdefault(photo.source.name, []).append(photo)
    matches: list[MatchedRow] = []
    errors: list[str] = []
    seen: Counter[UUID] = Counter()
    for row in rows:
        try:
            raw = row.values.get("file_name")
            if not isinstance(raw, str) or not raw:
                raise TableImportError(f"Row {row.line}: missing file_name")
            key = _file_key(raw)
            candidates = by_relative.get(key, []) if "/" in key else by_name.get(key, [])
            if not candidates:
                raise TableImportError(f"Row {row.line}: no photo matches {key}")
            if len(candidates) != 1:
                raise TableImportError(f"Row {row.line}: ambiguous photo {key}")
            photo = candidates[0]
            _validate_values(row, photo)
            seen[photo.id] += 1
            if seen[photo.id] > 1:
                raise TableImportError(f"Row {row.line}: duplicate photo {key}")
            matches.append(MatchedRow(photo.id, row))
        except TableImportError as error:
            errors.append(str(error))
    return ImportPreview(tuple(matches), tuple(errors))


def apply_manual_date(photo: PhotoItem, value: DateInput) -> PhotoItem:
    """Use the same date parser as imported rows while recording manual priority."""
    day = parse_capture_date(value)
    edits = dict(photo.edits)
    edits["capture_date_original"] = str(value)
    edits["capture_date_source"] = "manual"
    return replace(photo, edits=edits, taken_on=day)


def apply_import(preview: ImportPreview, photos: tuple[PhotoItem, ...]) -> tuple[PhotoItem, ...]:
    """Apply only an error-free preview; blank cells leave prior fields unchanged."""
    if preview.errors:
        raise TableImportError("Resolve preview errors before applying table edits")
    edits_by_id = {match.photo_id: match.row.values for match in preview.matches}
    updated: list[PhotoItem] = []
    for photo in photos:
        values = edits_by_id.get(photo.id)
        if values is None:
            updated.append(photo)
            continue
        edits = dict(photo.edits)
        for name in ("category", "subject", "location_name"):
            if name in values:
                edits[name] = str(values[name])
        day = photo.taken_on
        if "capture_date" in values:
            table_day = parse_capture_date(values["capture_date"])
            edits["imported_capture_date_original"] = str(values["capture_date"])
            if edits.get("capture_date_source") != "manual":
                day = table_day
                edits["capture_date_source"] = "table"
                edits["capture_date_original"] = str(values["capture_date"])
        coordinates = photo.coordinates
        if "latitude" in values:
            coordinates = (float(str(values["latitude"])), float(str(values["longitude"])))
        updated.append(replace(photo, edits=edits, taken_on=day, coordinates=coordinates))
    return tuple(updated)
